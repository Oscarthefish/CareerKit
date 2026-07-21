import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export(markdown_content: str, output_path: str | Path) -> Path:
    p = Path(output_path)
    p.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Default paragraph style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = markdown_content.split("\n")

    for line in lines:
        line_stripped = line.strip()

        if not line_stripped:
            doc.add_paragraph()
            continue

        if line_stripped.startswith("### "):
            heading = doc.add_heading(line_stripped[4:], level=3)
            heading.runs[0].font.color.rgb = RGBColor(0x1a, 0x56, 0x76)
        elif line_stripped.startswith("## "):
            heading = doc.add_heading(line_stripped[3:], level=2)
            heading.runs[0].font.color.rgb = RGBColor(0x0d, 0x3b, 0x52)
        elif line_stripped.startswith("# "):
            heading = doc.add_heading(line_stripped[2:], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line_stripped.startswith("- ") or line_stripped.startswith("* "):
            para = doc.add_paragraph(style="List Bullet")
            _add_inline(para, line_stripped[2:])
        elif line_stripped.startswith("**") and line_stripped.endswith("**"):
            para = doc.add_paragraph()
            run = para.add_run(line_stripped.strip("*"))
            run.bold = True
        elif re.match(r"^---+$", line_stripped):
            doc.add_paragraph("_" * 80)
        else:
            para = doc.add_paragraph()
            _add_inline(para, line_stripped)

    doc.save(str(p))
    return p


def _add_inline(para, text: str):
    """Parse inline bold/italic markdown and add runs to paragraph."""
    pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|(.+?)(?=\*|$))", re.DOTALL)
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*", text):
        if m.start() > pos:
            para.add_run(text[pos:m.start()])
        if m.group(0).startswith("**"):
            run = para.add_run(m.group(1))
            run.bold = True
        else:
            run = para.add_run(m.group(2))
            run.italic = True
        pos = m.end()
    if pos < len(text):
        para.add_run(text[pos:])
